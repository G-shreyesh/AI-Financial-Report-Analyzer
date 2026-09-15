import os
import pandas as pd

from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def read_pdf(file):
    reader = PdfReader(file)
    text = []

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text.append(page_text)

    return "\n\n".join(text)


def read_docx(file):
    document = Document(file)
    text = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells
            )

            if row_text.strip():
                text.append(row_text)

    return "\n".join(text)


def read_excel(file):
    excel_file = pd.ExcelFile(file)

    output = []

    for sheet_name in excel_file.sheet_names:
        df = pd.read_excel(
            file,
            sheet_name=sheet_name
        )

        output.append(
            f"\n--- SHEET: {sheet_name} ---\n"
        )

        if not df.empty:
            output.append(
                df.to_string(index=False)
            )

    return "\n".join(output)


def read_csv(file):
    df = pd.read_csv(file)

    return df.to_string(index=False)


def read_txt(file):
    file_bytes = file.getvalue()

    try:
        return file_bytes.decode("utf-8")

    except UnicodeDecodeError:
        return file_bytes.decode(
            "latin-1",
            errors="ignore"
        )


def read_pptx(file):
    presentation = Presentation(file)

    text = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        text.append(
            f"\n--- SLIDE {slide_number} ---\n"
        )

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                content = shape.text.strip()

                if content:
                    text.append(content)

    return "\n".join(text)


def extract_document_content(uploaded_file):
    """
    Convert supported company documents into text
    that can later be analyzed by the AI.
    """

    filename = uploaded_file.name

    extension = os.path.splitext(
        filename
    )[1].lower()

    try:

        if extension == ".pdf":

            content = read_pdf(
                uploaded_file
            )

        elif extension == ".docx":

            content = read_docx(
                uploaded_file
            )

        elif extension in [".xlsx", ".xls"]:

            content = read_excel(
                uploaded_file
            )

        elif extension == ".csv":

            content = read_csv(
                uploaded_file
            )

        elif extension == ".txt":

            content = read_txt(
                uploaded_file
            )

        elif extension == ".pptx":

            content = read_pptx(
                uploaded_file
            )

        else:

            return {
                "success": False,
                "filename": filename,
                "file_type": extension,
                "content": "",
                "error": (
                    f"Unsupported file type: "
                    f"{extension}"
                )
            }

        if not content.strip():

            return {
                "success": False,
                "filename": filename,
                "file_type": extension,
                "content": "",
                "error": (
                    "No readable text or data "
                    "was found in this file."
                )
            }

        return {
            "success": True,
            "filename": filename,
            "file_type": extension,
            "content": content,
            "error": None
        }

    except Exception as error:

        return {
            "success": False,
            "filename": filename,
            "file_type": extension,
            "content": "",
            "error": str(error)
        }